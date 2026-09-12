"""Generate SYNTHETIC interview coding (demo only) - 6 experts x 12 IQs, 11 final themes.
Updates filled workbook Interview_Coding + Pilot_Reliability sheets.
SYNTHETIC DEMONSTRATION ONLY - not real interviews.
"""
import os, json, pandas as pd, numpy as np
from datetime import date

BASE = os.path.dirname(os.path.abspath(__file__))
FILLED = os.path.join(BASE, "MIT_AI_Churn_SYNTHETIC_filled.xlsx")
rng = np.random.default_rng(7)

PARTICIPANTS = [
 ("P01", "Digital Banking Manager"),
 ("P02", "Data Scientist"),
 ("P03", "Digital Product Manager"),
 ("P04", "Data/AI Specialist"),
 ("P05", "Customer Experience Manager"),
 ("P06", "Product Manager"),
]
IQS = [f"IQ{i}" for i in range(1, 13)]
IQ_TOPICS = {
 "IQ1": "churn drivers observed in practice", "IQ2": "silent churn / multi-banking signs",
 "IQ3": "AI use for behavioural monitoring", "IQ4": "risk scoring and early warning",
 "IQ5": "explainability and reason codes", "IQ6": "retention targeting and prioritization",
 "IQ7": "treatment selection and timing", "IQ8": "human oversight and accountability",
 "IQ9": "data integration and quality", "IQ10": "cross-functional ownership",
 "IQ11": "governance, NDPA compliance, fairness", "IQ12": "what would prove value (pilot/metrics)",
}
THEMES = [
 ("Service quality", "Complaint resolution M=3.41", "Converges"),
 ("Trust", "Trust M=3.35", "Converges"),
 ("Behavioural signals", "Silent-churn detection M=3.18", "Converges"),
 ("Early warning", "Risk scoring M=3.38", "Converges"),
 ("Operational action", "AI insight integration M=2.58", "Complements"),
 ("Explainability", "Explainable AI M=3.14", "Complements"),
 ("Targeting", "Risk plus value M=3.40", "Converges"),
 ("Human oversight", "Human oversight M=3.18", "Converges"),
 ("Data integration", "Data integration M=2.70", "Complements"),
 ("Enterprise ownership", "Cross-functional ownership M=3.42", "Converges"),
 ("Service improvement", "Direct churn-reduction M=2.94", "Complements"),
]
QUOTES = {
 "Service quality": ["\"Repeated failures are what move customers, not one bad day. Fix the complaint loop first.\"",
   "\"Our churn cases almost always have an unresolved ticket behind them.\""],
 "Trust": ["\"Once trust drops after a fraud scare, activity migrates quietly.\"",
   "\"Customers stay where they feel money is safe and issues are owned.\""],
 "Behavioural signals": ["\"Declining logins and shrinking ticket size show up weeks before exit.\"",
   "\"Silent churn is the real problem; the account stays open but value leaves.\""],
 "Early warning": ["\"The model can identify a customer who is likely to leave, but if the bank does nothing with that information, the prediction has no commercial value.\" [P03 seed paraphrase - SYNTHETIC]",
   "\"Risk scores buy us time; without them we arrive after the customer has left.\""],
 "Operational action": ["\"Scores sitting in a dashboard change nothing; someone must own the next action.\"",
   "\"We need alert-to-action timing tracked like a service SLA.\""],
 "Explainability": ["\"Give me two or three reasons I can act on, not a black-box number.\"",
   "\"Reason codes decide whether I call, visit, or fix a fee.\""],
 "Targeting": ["\"Risk times value tells me who gets the scarce retention slot.\"",
   "\"Treat everyone the same and you waste budget and annoy customers.\""],
 "Human oversight": ["\"The data team can build the model, but they cannot be responsible for the entire retention outcome. The business has to own the outcome.\" [P04 seed paraphrase - SYNTHETIC]",
   "\"AI proposes, a named human disposes - especially on complaints.\""],
 "Data integration": ["\"Our signals live in five systems; stitching them is half the project.\"",
   "\"USSD versus app behaviour must be read differently.\""],
 "Enterprise ownership": ["\"Retention fails when it belongs to nobody; it needs business, data, tech and CX at one table.\"",
   "\"RACI and a weekly review beat another model tweak.\""],
 "Service improvement": ["\"No model compensates for persistent downtime - fix service, then optimize offers.\"",
   "\"Measure saved relationships, not just model AUC.\""],
}

rows = []
for pid, role in PARTICIPANTS:
    for iq in IQS:
        th, link, integ = THEMES[rng.integers(0, len(THEMES))]
        # force the two seed quotes for P03-IQ4 and P04-IQ8 for thesis continuity
        if pid == "P03" and iq == "IQ4":
            th, link, integ = "Early warning", "Risk scoring M=3.38", "Converges"
            q = QUOTES[th][0]
        elif pid == "P04" and iq == "IQ8":
            th, link, integ = "Human oversight", "Human oversight M=3.18", "Converges"
            q = QUOTES[th][0]
        else:
            q = QUOTES[th][rng.integers(0, 2)]
        rows.append({
         "Participant_ID": pid, "Role": role, "Interview_Date": f"2026-07-{rng.integers(1,28):02d}",
         "Question_Code": iq, "Initial_Code": f"{iq}: {IQ_TOPICS[iq]} - {th.lower()} noted",
         "Candidate_Theme": th, "Final_Theme": th,
         "Supporting_Quote": q + " [SYNTHETIC]",
         "Quant_Finding_Link": link, "Integration_Type": integ,
         "Analyst_Note": "SYNTHETIC demo coding for methods illustration."})
df = pd.DataFrame(rows)
df.to_csv(os.path.join(BASE, "interview_coding.csv"), index=False)
print(f"Wrote {len(df)} synthetic interview rows")

# update filled workbook
import openpyxl
wb = openpyxl.load_workbook(FILLED)
ws = wb["Interview_Coding"]
for r in range(2, 302):
    for c in range(1, 12):
        ws.cell(r, c).value = None
for i, (_, r) in enumerate(df.iterrows(), 2):
    for j, h in enumerate(df.columns, 1):
        ws.cell(i, j).value = r[h]
# Pilot_Reliability: fill observed alphas from reliability.csv
rel = pd.read_csv(os.path.join(BASE, "reliability.csv")).set_index("construct")
ws2 = wb["Pilot_Reliability"]
for row in ws2.iter_rows(min_row=2, max_row=8):
    con = row[0].value
    if con in rel.index:
        row[3].value = float(rel.loc[con, "pilot_alpha"])
        row[5].value = "Accept (alpha>=0.70, SYNTHETIC)" if rel.loc[con, "pilot_alpha"] >= 0.7 else "Review"
# Response_Tracker timestamp
ws3 = wb["Response_Tracker"]
ws3["B10"].value = pd.Timestamp.now()
wb.save(FILLED)
print("Updated filled workbook:", FILLED)
